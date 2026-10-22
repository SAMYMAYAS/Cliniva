"""Build compact offline WHO growth reference data from the official XLSX tables."""
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
Z_COLUMNS = {"-3": "SD3neg", "-2": "SD2neg", "-1": "SD1neg", "0": "SD0", "1": "SD1", "2": "SD2", "3": "SD3"}


def rows(filename, age_column, age_factor=1):
    sheet = load_workbook(ROOT / filename, read_only=True, data_only=True).active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    indexes = {name: headers.index(name) for name in Z_COLUMNS.values()}
    age_index = headers.index(age_column)
    result = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        age = row[age_index]
        if age is None:
            continue
        x = round(float(age) * age_factor, 2)
        result.append([x] + [round(float(row[indexes[name]]), 3) for name in Z_COLUMNS.values()])
    return result


def merge_age(a, b):
    # WHO 0-5 references use daily observations, while WHO 5-19 uses months.
    return a + [row for row in b if row[0] > 60]


data = {
    "source": "WHO Child Growth Standards (0–5 years) and WHO Growth Reference (5–19 years)",
    "z_levels": list(Z_COLUMNS),
    "age_unit": "month",
    "indicators": {
        "height_age": {
            "label": "Taille pour l’âge", "unit": "cm", "min_month": 0, "max_month": 228,
            "F": merge_age(rows("lhfa-girls-zscore-expanded-tables.xlsx", "Day", 1 / 30.4375), rows("hfa-5-19-girls.xlsx", "Month")),
            "M": merge_age(rows("hfa-boys-0-5.xlsx", "Day", 1 / 30.4375), rows("hfa-5-19-boys.xlsx", "Month")),
        },
        "weight_age": {
            "label": "Poids pour l’âge", "unit": "kg", "min_month": 0, "max_month": 120,
            "F": merge_age(rows("wfa-girls.xlsx", "Day", 1 / 30.4375), rows("wfa-5-10-girls.xlsx", "Month")),
            "M": merge_age(rows("wfa-boys.xlsx", "Day", 1 / 30.4375), rows("wfa-5-10-boys.xlsx", "Month")),
        },
        "bmi_age": {
            "label": "IMC pour l’âge", "unit": "kg/m²", "min_month": 0, "max_month": 228,
            "F": merge_age(rows("bfa-girls.xlsx", "Day", 1 / 30.4375), rows("bmi-5-19-girls.xlsx", "Month")),
            "M": merge_age(rows("bfa-boys.xlsx", "Day", 1 / 30.4375), rows("bmi-5-19-boys.xlsx", "Month")),
        },
        "head_age": {
            "label": "Périmètre crânien pour l’âge", "unit": "cm", "min_month": 0, "max_month": 60,
            "F": rows("hcfa-girls.xlsx", "Day", 1 / 30.4375), "M": rows("hcfa-boys.xlsx", "Day", 1 / 30.4375),
        },
        "arm_age": {
            "label": "Périmètre brachial pour l’âge", "unit": "cm", "min_month": 3, "max_month": 60,
            "F": rows("acfa-girls.xlsx", "Day", 1 / 30.4375), "M": rows("acfa-boys.xlsx", "Day", 1 / 30.4375),
        },
    },
}
(ROOT / "growth_reference.json").write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"Generated {ROOT / 'growth_reference.json'}")
